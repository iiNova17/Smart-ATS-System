"""Extract locally; original CV files are never sent to Kaggle."""

from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from ats.schemas import CVInput

MAX_FILE_BYTES = 10 * 1024 * 1024


def extract_cv(filename, data):
    if not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("File must be between 1 byte and 10 MB.")
    suffix = Path(filename).suffix.lower()
    stream = BytesIO(data)
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(stream)
        if reader.is_encrypted:
            raise ValueError("Password-protected PDF: upload an unlocked copy.")
        if len(reader.pages) > 30:
            raise ValueError("PDF exceeds 30 pages. Upload a CV of 30 pages or fewer.")
        texts = []
        for page in reader.pages:
            text = (page.extract_text() or "").strip()
            if len(text) < 10:
                raise ValueError(
                    "A PDF page has little readable text. OCR scanned pages first, then upload again."
                )
            texts.append(text)
        text = "\n\n".join(texts)
    elif suffix == ".docx":
        from docx import Document

        # Limit decompressed archive size before python-docx opens it.
        try:
            with ZipFile(stream) as archive:
                if sum(x.file_size for x in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError("DOCX expands beyond 40 MB; upload a smaller document.")
        except BadZipFile as exc:
            raise ValueError("This file is not a valid DOCX.") from exc
        stream.seek(0)
        document = Document(stream)
        # iter_inner_content preserves paragraph/table order for recent python-docx.
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        texts = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                texts.append(block.text)
            elif isinstance(block, Table):
                texts.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
        for section in document.sections:
            texts.extend(p.text for p in section.header.paragraphs)
            texts.extend(p.text for p in section.footer.paragraphs)
        text = "\n".join(texts)
    elif suffix == ".txt":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Save the TXT file with UTF-8 encoding and upload again.") from exc
    else:
        raise ValueError("Supported formats: PDF, DOCX, and TXT.")
    # Never silently truncate an applicant's CV.
    if len(text.strip()) < 40:
        raise ValueError("No usable CV text found. OCR scanned files before uploading.")
    if len(text) > 30000:
        raise ValueError("CV exceeds 30,000 characters. Upload a shorter version.")
    safe_name = filename.replace("\\", "/").split("/")[-1]
    return CVInput(filename=safe_name, text=text)
