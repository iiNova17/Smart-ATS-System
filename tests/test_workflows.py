import ast
import asyncio
import io
import json
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlparse

import pytest
import requests
from docx import Document
from fastapi.testclient import TestClient
from streamlit.testing.v1 import AppTest

from ats.client import APIClient, BackendError, api_url_from_env
from ats.extraction import extract_cv
from ats.matching import coverage_score, csv_export, skill_evidence
from ats.validation import validate_rank
from ats.workspace import (
    add_cv,
    analyze_cv,
    ask_cv,
    clear_cvs,
    embedding_response,
    get_cv,
    list_cvs,
    open_workspace,
    rank_cvs,
    remove_cv,
    search_cvs,
)
from backend.api import create_app
from backend.runtime import build_analysis_parser

ROOT = Path(__file__).resolve().parents[1]
# Synthetic text exists only inside automated tests; no candidate files ship with the app.
SAMPLE = """Test Candidate
Software Engineer
Skills: Python, SQL, PyTorch, NLP, LangChain, RAG, Streamlit, Git
Experience: Built Python document retrieval pipelines with LangChain and RAG.
Education: Computer Science. Project: Python document search with SQL storage.
"""
KEY = "a-long-enough-test-api-key-123456"


def make_test_runtime(state=None):
    state = state if state is not None else {"version": "test-embedding-v1"}

    def vector(text):
        return [1.0] + [float(skill in text.lower()) for skill in ["python", "java", "rag"]]

    def embed(texts):
        items = []
        for text in texts:
            chunks = [text[i : i + 120] for i in range(0, len(text), 120)]
            items.append(
                {
                    "vector": vector(text),
                    "chunks": [{"text": chunk, "vector": vector(chunk)} for chunk in chunks],
                }
            )
        return {"embedding_version": state["version"], "items": items}

    return {
        "embed_texts": embed,
        "answer_question": lambda context, question: context,
        "analyze_text": lambda text: {
            "name": "Test Candidate",
            "summary": "Profile from test model.",
            "skills": ["Python", "SQL"],
            "evidence": [
                {
                    "skill": "Python",
                    "quote": "Skills: Python, SQL, PyTorch, NLP, LangChain, RAG, Streamlit, Git",
                },
                {"skill": "Invented", "quote": "This does not appear in the CV."},
            ],
        },
    }


@pytest.fixture(autouse=True)
def isolated_data(monkeypatch, tmp_path):
    monkeypatch.setenv("ATS_DATA_DIR", str(tmp_path / "data"))


@pytest.fixture
def connected_workspace(monkeypatch, tmp_path):
    state = {"version": "test-embedding-v1"}
    server = TestClient(create_app(KEY, make_test_runtime(state)))
    calls = []

    def local_request(method, url, **kwargs):
        path = urlparse(url).path
        calls.append(
            {"path": path, "payload": kwargs.get("json"), "headers": kwargs.get("headers")}
        )
        result = server.request(
            method, path, json=kwargs.get("json"), headers=kwargs.get("headers")
        )
        response = requests.Response()
        response.status_code, response._content = result.status_code, result.content
        return response

    monkeypatch.setattr("ats.client.requests.request", local_request)
    monkeypatch.setenv("ATS_API_URL", "http://localhost:8000")
    monkeypatch.setenv("ATS_API_KEY", KEY)
    client = APIClient("http://localhost:8000", KEY)
    workspace = open_workspace(tmp_path / "data", lambda: client)
    return workspace, calls, state


def test_boundary_matching_and_punctuation():
    text = "Built C++ and C# services with .NET.\nUsed Python and machine\nlearning."
    for skill in ["C++", "C#", ".NET", "machine learning"]:
        assert skill_evidence(text, skill)
    assert skill_evidence("JavaScript and PostgreSQL", "Java") is None
    assert skill_evidence("JavaScript and PostgreSQL", "SQL") is None
    assert skill_evidence("RAGged rugs and Spark", "R") is None
    assert skill_evidence("Ｐｙｔｈｏｎ", "python") == "Ｐｙｔｈｏｎ"


def test_local_rank_deduplication_missing_evidence_and_weights(tmp_path):
    workspace = open_workspace(tmp_path / "data")
    add_cv(workspace, {"filename": "maya.txt", "text": SAMPLE})
    add_cv(
        workspace,
        {
            "filename": "other.txt",
            "text": "Candidate with Python and Docker. No other skills are listed in this CV.",
        },
    )
    rows = rank_cvs(
        workspace, {"required": ["Python", "SQL", " python "], "preferred": ["SQL", "RAG"]}
    )
    assert rows[0]["score"] == 100
    assert rows[1]["score"] == 40
    assert rows[1]["missing_required"] == ["SQL"]
    assert rows[1]["missing_preferred"] == ["RAG"]
    assert coverage_score(0, 2, 0, 1) == 50
    assert coverage_score(2, 0, 1, 0) == 50
    with pytest.raises(ValueError):
        validate_rank({"required": [" "]})


def test_offline_upload_search_rank_persist_and_delete_without_api(tmp_path):
    def forbidden_client():
        pytest.fail("A local operation contacted the model service.")

    directory = tmp_path / "data"
    workspace = open_workspace(directory, forbidden_client)
    cv = add_cv(workspace, {"filename": "first.txt", "text": SAMPLE})["cv"]
    assert add_cv(workspace, {"filename": "renamed.txt", "text": SAMPLE})["duplicate"]
    assert search_cvs(workspace, {"query": "Python, absent", "match": "all"}) == []
    assert len(search_cvs(workspace, {"query": "Python, absent", "match": "any"})) == 1
    assert rank_cvs(workspace, {"required": ["Python"]})[0]["score"] == 100
    reloaded = open_workspace(directory, forbidden_client)
    assert get_cv(reloaded, cv["id"])["text"] == SAMPLE.strip()
    remove_cv(reloaded, cv["id"])
    assert list_cvs(workspace) == []
    with pytest.raises(KeyError):
        get_cv(workspace, cv["id"])


def test_local_library_has_no_fixed_candidate_or_search_cap(tmp_path):
    workspace = open_workspace(tmp_path)
    for i in range(105):
        add_cv(
            workspace, {"filename": f"cv_{i}.txt", "text": SAMPLE + f"\nCandidate reference {i}."}
        )
    assert len(list_cvs(workspace)) == 105
    assert len(search_cvs(workspace, {"query": "Python"})) == 105
    assert len(rank_cvs(workspace, {"required": ["Python"]})) == 105
    clear_cvs(workspace)
    assert list_cvs(workspace) == []


def test_extraction_docx_table_order_and_utf8():
    document = Document()
    document.add_paragraph("Candidate with professional experience and projects.")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Skills: Python, SQL"
    document.add_paragraph("Education: BSc Computer Science")
    buffer = io.BytesIO()
    document.save(buffer)
    cv = extract_cv("../../candidate.docx", buffer.getvalue())
    assert cv.filename == "candidate.docx"
    assert cv.text.index("Skills:") < cv.text.index("Education:")
    assert extract_cv("cv.txt", SAMPLE.encode()).text == SAMPLE.strip()
    with pytest.raises(ValueError, match="encoding"):
        extract_cv("cv.txt", b"\xff" * 100)
    with pytest.raises(ValueError, match="30,000"):
        extract_cv("long.txt", b"a" * 30001)


def test_scanned_and_encrypted_pdf_messages():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match="OCR"):
        extract_cv("scan.pdf", buffer.getvalue())
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match="Password"):
        extract_cv("locked.pdf", buffer.getvalue())


def test_model_api_is_authenticated_stateless_and_inference_only():
    client = TestClient(create_app(KEY, make_test_runtime()))
    headers = {"Authorization": "Bearer " + KEY}
    assert client.get("/health").status_code == 401
    assert client.get("/health", headers=headers).json()["version"] == "3.0.0"
    assert client.get("/cvs", headers=headers).status_code == 404
    assert client.post("/search", headers=headers, json={"query": "Python"}).status_code == 404
    assert client.post("/rank", headers=headers, json={"required": ["Python"]}).status_code == 404
    assert client.post("/analyze", headers=headers, json={"text": SAMPLE}).json()["analysis"][
        "name"
    ]
    assert client.post("/embed", headers=headers, json={"texts": ["Python"]}).json()["items"]
    answer = client.post(
        "/answer",
        headers=headers,
        json={"question": "Skills?", "passages": ["Skills: Python, SQL."]},
    ).json()
    assert "[Passage 1]" in answer["answer"]
    for path, payload in [
        ("/embed", {"texts": []}),
        ("/embed", {"texts": ["a"] * 9}),
        ("/analyze", {"text": "short"}),
        ("/answer", {"question": "Skills?", "passages": []}),
    ]:
        assert client.post(path, headers=headers, json=payload).status_code == 422


def test_api_model_failure_omits_internal_details():
    def broken(text):
        raise RuntimeError("Private runtime information")

    client = TestClient(create_app(KEY, {"analyze_text": broken}))
    response = client.post(
        "/analyze", headers={"Authorization": "Bearer " + KEY}, json={"text": SAMPLE}
    )
    assert response.status_code == 503
    assert "Private" not in response.text


def test_profiles_are_grounded_cached_and_available_after_restart(connected_workspace):
    workspace, calls, state = connected_workspace
    cv_id = add_cv(workspace, {"filename": "maya.txt", "text": SAMPLE})["cv"]["id"]
    assert calls == []
    profile = analyze_cv(workspace, cv_id)
    assert len(profile["evidence"]) == 1
    assert profile["review_notes"]
    assert calls[-1]["path"] == "/analyze"
    assert calls[-1]["payload"] == {"text": SAMPLE.strip()}
    restarted_offline = open_workspace(workspace["database"].parent)
    assert analyze_cv(restarted_offline, cv_id) == profile
    assert len(calls) == 1
    assert list_cvs(restarted_offline)[0]["name"] == "Test Candidate"
    clear_cvs(restarted_offline)
    assert list_cvs(workspace) == []


def test_json_output_parser_rejects_truncation_and_wrong_types():
    from langchain_core.exceptions import OutputParserException

    parser = build_analysis_parser()
    assert parser.invoke('{"name": "Alex", "skills": ["Python"]}')["name"] == "Alex"
    with pytest.raises(OutputParserException):
        parser.invoke('{"name": "Alex", "skills": ["Python"')
    with pytest.raises(ValueError):
        parser.invoke('{"experience": "wrong type"}')
    with pytest.raises(ValueError):
        parser.invoke('{"experience": [{"details": [42]}]}')


def test_similarity_and_rag_are_local_and_cached(connected_workspace):
    workspace, calls, state = connected_workspace
    cv_id = add_cv(workspace, {"filename": "maya.txt", "text": SAMPLE})["cv"]["id"]
    second = "Different candidate. Skills: Java, Docker. Built Java services and Java projects."
    add_cv(workspace, {"filename": "second.txt", "text": second})
    results = search_cvs(workspace, {"query": "Python", "mode": "semantic"})
    assert len(results) == 2 and results[0]["id"] == cv_id
    assert [call["path"] for call in calls] == ["/embed", "/embed"]
    count = len(calls)
    search_cvs(workspace, {"query": "Python", "mode": "semantic"})
    assert len(calls) == count + 1  # Only the new query is embedded.
    answer = ask_cv(workspace, cv_id, "Python projects?")
    assert all("Different candidate" not in passage for passage in answer["passages"])
    assert calls[-1]["path"] == "/answer"
    assert set(calls[-1]["payload"]) == {"question", "passages"}
    assert calls[-1]["payload"]["passages"] == answer["passages"]
    assert all("X-Session-ID" not in call["headers"] for call in calls)
    remove_cv(workspace, cv_id)
    assert all(
        row["id"] != cv_id for row in search_cvs(workspace, {"query": "Python", "mode": "semantic"})
    )


def test_embedding_model_change_rebuilds_local_cache(connected_workspace):
    workspace, calls, state = connected_workspace
    add_cv(workspace, {"filename": "maya.txt", "text": SAMPLE})
    search_cvs(workspace, {"query": "Python", "mode": "semantic"})
    state["version"] = "test-embedding-v2"
    calls.clear()
    search_cvs(workspace, {"query": "Python", "mode": "semantic"})
    assert len(calls) == 2
    assert calls[1]["payload"]["texts"] == [SAMPLE.strip()]


def test_embeddings_reject_unsupported_passages_and_invalid_vectors():
    response = {
        "embedding_version": "v1",
        "items": [{"vector": [1.0], "chunks": [{"text": "invented passage", "vector": [1.0]}]}],
    }
    with pytest.raises(BackendError, match="unsupported"):
        embedding_response(response, ["actual passage"])
    response["items"][0]["vector"] = [float("nan")]
    with pytest.raises(BackendError, match="invalid"):
        embedding_response(response, ["actual passage"])


def test_description_ranking_calls_embeddings_only(connected_workspace):
    workspace, calls, state = connected_workspace
    add_cv(workspace, {"filename": "maya.txt", "text": SAMPLE})
    rows = rank_cvs(
        workspace, {"required": ["Python", "SQL"], "job_description": "Python and RAG projects"}
    )
    assert rows[0]["score"] == 100 and rows[0]["semantic_similarity"] is not None
    assert all(call["path"] == "/embed" for call in calls)


def test_csv_export_neutralizes_formulas():
    csv = csv_export([{"Document": '=HYPERLINK("bad")', "Quote": "  @bad"}])
    assert "'=HYPERLINK" in csv and "'  @bad" in csv


def test_client_uses_developers_domain_and_optional_full_url_override(monkeypatch):
    monkeypatch.delenv("ATS_API_URL", raising=False)
    monkeypatch.setenv("NGROK_DOMAIN", "owned-domain.ngrok-free.dev")
    assert api_url_from_env() == "https://owned-domain.ngrok-free.dev"
    client = APIClient(api_url_from_env(), KEY)
    assert client.url == "https://owned-domain.ngrok-free.dev"
    assert client.headers["ngrok-skip-browser-warning"] == "true"
    monkeypatch.setenv("ATS_API_URL", "http://localhost:8000")
    assert api_url_from_env() == "http://localhost:8000"
    monkeypatch.delenv("ATS_API_URL")
    monkeypatch.setenv("NGROK_DOMAIN", "https://invalid.ngrok-free.dev")
    with pytest.raises(BackendError):
        api_url_from_env()
    monkeypatch.delenv("NGROK_DOMAIN")
    assert api_url_from_env() == ""


def test_notebook_sdk_uses_account_configuration_and_closes_listener(monkeypatch, capsys):
    import sys

    notebook = json.loads((ROOT / "notebooks/smart_ats_kaggle.ipynb").read_text(encoding="utf-8"))
    blocks = ["".join(cell["source"]) for cell in notebook["cells"] if cell["cell_type"] == "code"]
    settings = {
        "NGROK_DOMAIN": "owned-domain.ngrok-free.dev",
        "NGROK_AUTHTOKEN": "private-token",
        "ATS_API_KEY": KEY,
    }
    for name, value in settings.items():
        monkeypatch.setenv(name, value)
    fake_secrets = SimpleNamespace(
        UserSecretsClient=lambda: SimpleNamespace(get_secret=settings.get)
    )
    monkeypatch.setitem(sys.modules, "kaggle_secrets", fake_secrets)
    namespace = {}
    exec(blocks[1], namespace)
    calls, closed = [], []

    async def close():
        closed.append(True)

    async def forward(address, **kwargs):
        calls.append((address, kwargs))
        return SimpleNamespace(url=lambda: "https://" + settings["NGROK_DOMAIN"], close=close)

    monkeypatch.setitem(sys.modules, "ngrok", SimpleNamespace(forward=forward))

    async def run_cells():
        for index in [7, 7, 9]:
            compiled = compile(
                blocks[index], "sdk_cell", "exec", flags=ast.PyCF_ALLOW_TOP_LEVEL_AWAIT
            )
            await eval(compiled, namespace)

    asyncio.run(run_cells())
    assert (
        calls
        == [("localhost:8000", {"authtoken_from_env": True, "domain": settings["NGROK_DOMAIN"]})]
        * 2
    )
    assert len(closed) == 2
    assert "ats_forwarder" not in namespace
    output = capsys.readouterr().out
    assert settings["NGROK_AUTHTOKEN"] not in output and KEY not in output


def test_connection_check_reports_health_without_cv_or_credentials(connected_workspace, capsys):
    from tools.check_connection import main

    workspace, calls, state = connected_workspace
    assert main() == 0
    assert [call["path"] for call in calls] == ["/health"]
    output = capsys.readouterr().out
    assert "Connected" in output and KEY not in output


def test_product_ui_upload_library_search_rank_do_not_contact_models(connected_workspace):
    workspace, calls, state = connected_workspace
    add_cv(workspace, extract_cv("maya.txt", SAMPLE.encode()).model_dump())
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=20).run()
    assert not app.exception and calls == []
    assert app.metric[0].value == "1"
    assert not app.radio
    assert not any(widget.label in {"API URL", "API key", "Backend"} for widget in app.text_input)
    assert all("Capacity" not in metric.label for metric in app.metric)
    app.button(key="nav_search_candidates").click().run()
    next(
        widget for widget in app.text_input if widget.label == "What are you looking for?"
    ).set_value("Python, LangChain")
    next(button for button in app.button if button.label == "Find candidates").click().run()
    assert len(app.session_state["search_results"]["rows"]) == 1
    app.button(key="nav_rank_candidates").click().run()
    next(widget for widget in app.text_area if widget.label == "Required skills").set_value(
        "Python, SQL"
    )
    next(button for button in app.button if button.label == "Create shortlist").click().run()
    assert not app.exception
    assert app.session_state["rank_results"]["rows"][0]["score"] == 100
    assert calls == []


def test_product_profile_and_question_form(connected_workspace):
    workspace, calls, state = connected_workspace
    add_cv(workspace, {"filename": "maya.txt", "text": SAMPLE})
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=20).run()
    app.button(key="nav_cv_analysis").click().run()
    next(button for button in app.button if button.label == "Prepare profile").click().run()
    assert not app.exception
    next(widget for widget in app.text_input if widget.label == "Your question").set_value(
        "Which Python skills are mentioned?"
    )
    next(button for button in app.button if button.label == "Ask question").click().run()
    assert not app.exception
    assert "[Passage 1]" in next(iter(app.session_state["answers"].values()))["answer"]


def test_unconfigured_product_stays_usable_and_reports_only_ai_failure(monkeypatch):
    monkeypatch.delenv("ATS_API_URL", raising=False)
    monkeypatch.delenv("ATS_API_KEY", raising=False)
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=20).run()
    assert not app.exception and not app.radio
    assert not any("Demo" in b.label or "Connect" in b.label for b in app.button)
    assert not any(t.label in {"API URL", "API key"} for t in app.text_input)
    assert not app.info and not app.error
    assert app.metric[0].value == "0"


def test_notebook_is_standalone_compiles_and_contains_only_inference():
    notebook = json.loads((ROOT / "notebooks/smart_ats_kaggle.ipynb").read_text(encoding="utf-8"))
    blocks = ["".join(cell["source"]) for cell in notebook["cells"] if cell["cell_type"] == "code"]
    assert len(blocks) == 10
    for i, source in enumerate(blocks):
        if source.startswith("%pip"):
            continue
        tree = ast.parse(source)
        assert not any(isinstance(node, ast.ClassDef) for node in ast.walk(tree))
        assert "from ats." not in source and "from backend." not in source
        compile(source, f"block_{i + 1}", "exec", flags=ast.PyCF_ALLOW_TOP_LEVEL_AWAIT)
    namespace = {}
    exec(blocks[2], namespace)
    exec(blocks[5], namespace)
    client = TestClient(namespace["create_app"](KEY, make_test_runtime()))
    headers = {"Authorization": "Bearer " + KEY}
    assert client.post("/embed", headers=headers, json={"texts": ["Python"]}).status_code == 200
    assert client.post("/analyze", headers=headers, json={"text": SAMPLE}).status_code == 200
    assert (
        client.post(
            "/answer", headers=headers, json={"question": "Skills?", "passages": ["Python"]}
        ).status_code
        == 200
    )
    assert client.get("/cvs", headers=headers).status_code == 404
